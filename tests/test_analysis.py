"""
Tests for the layout manifest system (`src/core/analysis/`):
structural pass, models, LLM-classification merge, caching, and CLI wiring.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest

from src.core.analysis.models import SlideManifest, SlotManifest, TemplateManifest
from src.core.analysis.structural import (
    build_docx_section_manifests,
    build_pptx_slide_manifest,
    classify_pptx_slide_structure,
)
from src.core.models import DocumentSection, ParsedDocument, SlideContent, SlideIntent
from src.core.planner.content_planner import ContentPlanner

PPTX_TEMPLATE = Path(__file__).parent.parent / "templates" / "generic-slides.pptx"
DOCX_TEMPLATE = Path(__file__).parent.parent / "templates" / "generic-document.docx"


# ---------------------------------------------------------------------------
# Structural pass — PPTX
# ---------------------------------------------------------------------------

def test_classify_pptx_slide_structure_key_highlights():
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from pptx import Presentation

    prs = Presentation(str(PPTX_TEMPLATE))
    signals = classify_pptx_slide_structure(prs.slides[7])
    assert signals["guessed_content_affinity"] in ("key_highlights", "grid", "features_or_benefits")
    assert signals["short_text_shapes"] >= 3


def test_build_pptx_slide_manifest_key_highlights_has_card_slots():
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from pptx import Presentation

    prs = Presentation(str(PPTX_TEMPLATE))
    manifest = build_pptx_slide_manifest(7, prs.slides[7], prs.slide_width, prs.slide_height)
    assert manifest.index == 7
    assert any(slot.injection_recipe == "apply_items" for slot in manifest.slots)
    assert manifest.capacity  # width/height were supplied


def test_build_pptx_slide_manifest_bullets_slide_has_bullets_slot():
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from pptx import Presentation

    prs = Presentation(str(PPTX_TEMPLATE))
    manifest = build_pptx_slide_manifest(3, prs.slides[3], prs.slide_width, prs.slide_height)
    assert any(slot.injection_recipe == "set_bullets" for slot in manifest.slots)


# ---------------------------------------------------------------------------
# Structural pass — DOCX
# ---------------------------------------------------------------------------

def test_build_docx_section_manifests_detects_body_section():
    if not DOCX_TEMPLATE.exists():
        pytest.skip("Bundled DOCX template not available")
    from docx import Document

    doc = Document(str(DOCX_TEMPLATE))
    sections = build_docx_section_manifests(doc)
    roles = [s.role for s in sections]
    assert "body" in roles
    body = next(s for s in sections if s.role == "body")
    assert isinstance(body.supports_table, bool)
    assert isinstance(body.has_image_anchor, bool)


# ---------------------------------------------------------------------------
# Models — serialization, best_slide_for
# ---------------------------------------------------------------------------

def test_template_manifest_json_roundtrip(tmp_path):
    manifest = TemplateManifest(
        template_path="foo.pptx",
        format="pptx",
        generated_at=TemplateManifest.timestamp(),
        slides=[
            SlideManifest(
                index=7,
                layout_name="Key Highlights",
                content_affinity=["key_highlights"],
                capacity={"max_items": 4},
                slots=[SlotManifest(role="card_1", injection_recipe="apply_items", accepts=["text"])],
            ),
        ],
    )
    out = manifest.save(tmp_path / "m.json")
    loaded = TemplateManifest.load(out)
    assert loaded.format == "pptx"
    assert len(loaded.slides) == 1
    assert loaded.slides[0].content_affinity == ["key_highlights"]
    assert loaded.slides[0].slots[0].injection_recipe == "apply_items"


def test_best_slide_for_matches_affinity_and_capacity():
    manifest = TemplateManifest(slides=[
        SlideManifest(index=3, content_affinity=["bullets"], capacity={"body_max_chars_per_slot": 200}),
        SlideManifest(index=7, content_affinity=["key_highlights"], capacity={"max_items": 4}),
    ])
    best = manifest.best_slide_for(["bullets"])
    assert best is not None
    assert best.index == 3

    assert manifest.best_slide_for(["nonexistent_affinity"]) is None


def test_best_slide_for_requires_image_slot():
    manifest = TemplateManifest(slides=[
        SlideManifest(index=3, content_affinity=["diagram"], slots=[]),
        SlideManifest(
            index=5, content_affinity=["diagram"],
            slots=[SlotManifest(role="image_slot_1", accepts=["image"])],
        ),
    ])
    best = manifest.best_slide_for(["diagram"], needs_image=True)
    assert best is not None
    assert best.index == 5


# ---------------------------------------------------------------------------
# manifest.py — generation, merge, caching
# ---------------------------------------------------------------------------

def test_generate_pptx_manifest_structural_only():
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from src.core.analysis.manifest import generate_pptx_manifest

    manifest = generate_pptx_manifest(PPTX_TEMPLATE, llm_provider=None)
    assert manifest.format == "pptx"
    assert manifest.classified_by_llm is False
    assert len(manifest.slides) == 13
    assert manifest.dimensions["slide_width_emu"] > 0


def test_generate_docx_manifest_structural_only():
    if not DOCX_TEMPLATE.exists():
        pytest.skip("Bundled DOCX template not available")
    from src.core.analysis.manifest import generate_docx_manifest

    manifest = generate_docx_manifest(DOCX_TEMPLATE, llm_provider=None)
    assert manifest.format == "docx"
    assert manifest.classified_by_llm is False
    assert any(s.role == "body" for s in manifest.sections)


def test_generate_pptx_manifest_with_llm_merges_affinity_and_validates_recipes():
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from src.core.analysis.manifest import generate_pptx_manifest
    from src.core.llm.base import LLMProvider

    class FakeProvider(LLMProvider):
        def generate(self, prompt: str) -> str:
            return json.dumps({"content_affinity": ["totally_made_up_label"], "notes": "a test slide"})

    manifest = generate_pptx_manifest(PPTX_TEMPLATE, llm_provider=FakeProvider())
    assert manifest.classified_by_llm is True
    slide0 = manifest.slides[0]
    # The LLM's made-up affinity label is still accepted (affinity is a free
    # label, not restricted to the fixed enum — only injection_recipe is).
    assert slide0.content_affinity == ["totally_made_up_label"]
    assert "a test slide" in slide0.notes


def test_generate_pptx_manifest_llm_failure_falls_back_to_structural():
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from src.core.analysis.manifest import generate_pptx_manifest
    from src.core.llm.base import LLMProvider

    class BrokenProvider(LLMProvider):
        def generate(self, prompt: str) -> str:
            raise RuntimeError("provider unreachable")

    manifest = generate_pptx_manifest(PPTX_TEMPLATE, llm_provider=BrokenProvider())
    assert manifest.classified_by_llm is False
    assert len(manifest.slides) == 13


def test_merge_never_accepts_invalid_injection_recipe_via_slots():
    """The manifest's `INJECTION_RECIPES` enum is the safety boundary: an
    LLM can never cause a slot's `injection_recipe` to become something not
    in that fixed set, since `_merge_slide` never touches `slots` at all."""
    from src.core.analysis.manifest import _merge_slide
    from src.core.analysis.models import INJECTION_RECIPES

    slide = SlideManifest(
        index=0,
        slots=[SlotManifest(role="title", injection_recipe="single_run_replace")],
    )
    merged = _merge_slide(slide, {"content_affinity": ["x"], "injection_recipe_override": "rm -rf /"})
    assert merged.slots[0].injection_recipe == "single_run_replace"
    assert merged.slots[0].injection_recipe in INJECTION_RECIPES


def test_get_or_generate_manifest_caches_and_reuses(tmp_path):
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    import shutil

    from src.core.analysis.manifest import get_or_generate_manifest, manifest_cache_path

    local_template = tmp_path / "t.pptx"
    shutil.copy(PPTX_TEMPLATE, local_template)

    manifest1 = get_or_generate_manifest(local_template)
    cache_path = manifest_cache_path(local_template)
    assert cache_path.exists()
    mtime1 = cache_path.stat().st_mtime

    manifest2 = get_or_generate_manifest(local_template)
    assert cache_path.stat().st_mtime == mtime1  # reused cache, not regenerated
    assert len(manifest2.slides) == len(manifest1.slides)

    manifest3 = get_or_generate_manifest(local_template, force=True)
    assert len(manifest3.slides) == len(manifest1.slides)


def test_get_or_generate_manifest_rejects_xlsx(tmp_path):
    fake_xlsx = tmp_path / "t.xlsx"
    fake_xlsx.write_bytes(b"not a real xlsx")
    from src.core.analysis.manifest import get_or_generate_manifest

    with pytest.raises(ValueError, match="not supported"):
        get_or_generate_manifest(fake_xlsx)


# ---------------------------------------------------------------------------
# CLI wiring
# ---------------------------------------------------------------------------

def test_cli_analyze_template_manifest(tmp_path):
    if not PPTX_TEMPLATE.exists():
        pytest.skip("Bundled PPTX template not available")
    from typer.testing import CliRunner

    from src.cli.main import app
    from src.core.analysis.manifest import manifest_cache_path

    out_path = tmp_path / "manifest.json"
    runner = CliRunner()
    result = runner.invoke(app, [
        "analyze-template", str(PPTX_TEMPLATE),
        "--manifest", str(out_path),
        "--llm", "none",
    ])
    assert result.exit_code == 0, result.output
    assert out_path.exists()
    data = json.loads(out_path.read_text())
    assert data["format"] == "pptx"
    assert len(data["slides"]) == 13
    # Regression: --manifest must write ONLY the path the caller named, not
    # also a second copy at the bundled template's own cache location.
    assert not manifest_cache_path(PPTX_TEMPLATE).exists()


# ---------------------------------------------------------------------------
# ContentPlanner integration
# ---------------------------------------------------------------------------

def test_content_planner_uses_manifest_when_supplied():
    manifest = TemplateManifest(slides=[
        SlideManifest(index=42, content_affinity=["bullets"], capacity={}),
    ])
    planner = ContentPlanner(manifest=manifest)

    document = ParsedDocument(
        title="Doc", source_path=Path("x.md"), source_format="md",
        sections=[DocumentSection(title="Section", level=1)],
    )
    normalized = [(SlideIntent.BULLETS, SlideContent(title="Section", bullets=["a", "b"]))]

    plan = planner.plan_slides(document, normalized)
    assert plan.slides[0].template_index == 42


def test_content_planner_falls_back_to_static_map_without_affinity_match():
    manifest = TemplateManifest(slides=[
        SlideManifest(index=42, content_affinity=["some_other_affinity"], capacity={}),
    ])
    planner = ContentPlanner(manifest=manifest)

    document = ParsedDocument(
        title="Doc", source_path=Path("x.md"), source_format="md",
        sections=[DocumentSection(title="Section", level=1)],
    )
    normalized = [(SlideIntent.BULLETS, SlideContent(title="Section", bullets=["a"]))]

    plan = planner.plan_slides(document, normalized)
    # No manifest slide has a "bullets" affinity, so falls back to the
    # static DEFAULT_TEMPLATE_MAP's bullets index (3).
    assert plan.slides[0].template_index == 3


def test_content_planner_without_manifest_is_unchanged():
    """Behavior for every existing caller (no `manifest` argument) must be
    byte-identical to before layout manifests existed."""
    planner = ContentPlanner()
    document = ParsedDocument(
        title="Doc", source_path=Path("x.md"), source_format="md",
        sections=[DocumentSection(title="Section", level=1)],
    )
    normalized = [(SlideIntent.BULLETS, SlideContent(title="Section", bullets=["a"]))]
    plan = planner.plan_slides(document, normalized)
    assert plan.slides[0].template_index == 3
