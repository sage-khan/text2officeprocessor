"""
Tests for the XLSX Engine.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from src.core.engines.xlsx.engine import XLSXEngine
from src.core.exceptions import RenderError
from src.core.models import SheetDefinition, SpreadsheetPlan
from src.core.parser.preprocessor import InputPreprocessor
from src.core.planner.content_planner import ContentPlanner


@pytest.fixture
def engine():
    return XLSXEngine()


@pytest.fixture
def simple_plan():
    return SpreadsheetPlan(
        sheets=[
            SheetDefinition(
                name="Metrics",
                columns=["Metric", "Value", "Unit"],
                rows=[
                    ["Accuracy", "0.95", "ratio"],
                    ["Precision", "0.88", "ratio"],
                    ["Recall", "0.91", "ratio"],
                ],
            )
        ]
    )


def test_render_produces_file(engine, simple_plan, tmp_path):
    out = tmp_path / "output.xlsx"
    result = engine.render(simple_plan, out)
    assert result.exists()


def test_output_is_valid_xlsx(engine, simple_plan, tmp_path):
    out = tmp_path / "output.xlsx"
    engine.render(simple_plan, out)
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    assert "Metrics" in wb.sheetnames


def test_header_row_is_written(engine, simple_plan, tmp_path):
    out = tmp_path / "output.xlsx"
    engine.render(simple_plan, out)
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    ws = wb["Metrics"]
    assert ws.cell(row=1, column=1).value == "Metric"
    assert ws.cell(row=1, column=2).value == "Value"


def test_data_rows_written(engine, simple_plan, tmp_path):
    out = tmp_path / "output.xlsx"
    engine.render(simple_plan, out)
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    ws = wb["Metrics"]
    assert ws.cell(row=2, column=1).value == "Accuracy"
    assert ws.cell(row=4, column=2).value == "0.91"


def test_empty_plan_raises_error(engine, tmp_path):
    with pytest.raises(RenderError, match="no sheets"):
        engine.render(SpreadsheetPlan(sheets=[]), tmp_path / "out.xlsx")


def test_multi_sheet_plan(engine, tmp_path):
    plan = SpreadsheetPlan(
        sheets=[
            SheetDefinition(name="Sheet1", columns=["A"], rows=[["1"], ["2"]]),
            SheetDefinition(name="Sheet2", columns=["X", "Y"], rows=[["a", "b"]]),
        ]
    )
    out = tmp_path / "multi.xlsx"
    engine.render(plan, out)
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    assert set(wb.sheetnames) == {"Sheet1", "Sheet2"}


def test_from_markdown_table(tmp_path):
    """Full pipeline: md with table → SpreadsheetPlan → xlsx."""
    f = tmp_path / "report.md"
    f.write_text(
        "# Report\n\n## Bias Metrics\n\n"
        "| Group | FPR | FNR |\n|-------|-----|-----|\n"
        "| Black | 0.45 | 0.28 |\n| White | 0.23 | 0.31 |\n",
        encoding="utf-8",
    )
    parsed = InputPreprocessor().parse(f)
    plan = ContentPlanner.plan_spreadsheet(parsed)
    out = tmp_path / "output.xlsx"
    XLSXEngine().render(plan, out)
    assert out.exists()
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    assert len(wb.sheetnames) >= 1


def test_sheet_name_is_sanitized_and_unique(engine, tmp_path):
    plan = SpreadsheetPlan(
        sheets=[
            SheetDefinition(name="Revenue: Q1/Q2*?", columns=["A"], rows=[["1"]]),
            SheetDefinition(name="Revenue: Q1/Q2*?", columns=["A"], rows=[["2"]]),
        ]
    )
    out = tmp_path / "sanitized.xlsx"
    engine.render(plan, out)
    import openpyxl

    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames[0] == "Revenue_ Q1_Q2__"
    assert wb.sheetnames[1] == "Revenue_ Q1_Q2___1"
