"""
Tests for draw.io converter and diagram embedding in the PPTX pipeline.
"""
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from src.core.drawio.converter import (
    DrawioExportError,
    _detect_page_count,
    export_drawio_to_png,
)
from src.core.models import SlideDefinition, SlideIntent
from src.core.planner.content_planner import ContentPlanner


# ---------------------------------------------------------------------------
# Converter unit tests (subprocess mocked — no actual drawio CLI required)
# ---------------------------------------------------------------------------

def test_export_raises_if_file_missing(tmp_path):
    with pytest.raises(FileNotFoundError):
        export_drawio_to_png(tmp_path / "missing.drawio")


def test_export_raises_if_drawio_not_on_path(tmp_path):
    src = tmp_path / "diagram.drawio"
    src.write_text("<mxfile><diagram></diagram></mxfile>", encoding="utf-8")
    with patch("src.core.drawio.converter._drawio_available", return_value=False):
        with pytest.raises(DrawioExportError, match="drawio CLI not found"):
            export_drawio_to_png(src)


def test_export_raises_on_nonzero_exit(tmp_path):
    src = tmp_path / "diagram.drawio"
    src.write_text("<mxfile><diagram></diagram></mxfile>", encoding="utf-8")

    mock_result = MagicMock()
    mock_result.returncode = 1
    mock_result.stderr = "Something went wrong"

    with patch("src.core.drawio.converter._drawio_available", return_value=True), \
         patch("src.core.drawio.converter._xvfb_available", return_value=False), \
         patch("subprocess.run", return_value=mock_result):
        with pytest.raises(DrawioExportError, match="exit 1"):
            export_drawio_to_png(src, output_path=tmp_path / "out.png")


def test_export_succeeds_and_returns_path(tmp_path):
    src = tmp_path / "diagram.drawio"
    src.write_text("<mxfile><diagram></diagram></mxfile>", encoding="utf-8")
    out = tmp_path / "out.png"
    out.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)  # minimal fake PNG

    mock_result = MagicMock()
    mock_result.returncode = 0

    with patch("src.core.drawio.converter._drawio_available", return_value=True), \
         patch("src.core.drawio.converter._xvfb_available", return_value=False), \
         patch("subprocess.run", return_value=mock_result):
        result = export_drawio_to_png(src, output_path=out)

    assert result == out


def test_detect_page_count_single(tmp_path):
    f = tmp_path / "one.drawio"
    f.write_text("<mxfile><diagram id='1'>data</diagram></mxfile>", encoding="utf-8")
    assert _detect_page_count(f) == 1


def test_detect_page_count_multi(tmp_path):
    f = tmp_path / "multi.drawio"
    f.write_text(
        "<mxfile>"
        "<diagram id='1'>data</diagram>"
        "<diagram id='2'>data</diagram>"
        "<diagram id='3'>data</diagram>"
        "</mxfile>",
        encoding="utf-8",
    )
    assert _detect_page_count(f) == 3


# ---------------------------------------------------------------------------
# Planner: - diagram: line parsing
# ---------------------------------------------------------------------------

def test_planner_parses_diagram_line(tmp_path):
    slides_md = tmp_path / "slides.md"
    slides_md.write_text(
        '## SLIDE 1 — template_index: 2 (Single Point)\n'
        '- placeholder: "SINGLE POINT SLIDE" → "Architecture Overview"\n'
        '- diagram: "diagrams/architecture.drawio"\n',
        encoding="utf-8",
    )
    plan = ContentPlanner.parse_slides_markdown(slides_md)
    assert len(plan.slides) == 1
    slide = plan.slides[0]
    assert slide.diagram_path == "diagrams/architecture.drawio"
    assert slide.intent == SlideIntent.DIAGRAM


def test_planner_no_diagram_line_leaves_empty(tmp_path):
    slides_md = tmp_path / "slides.md"
    slides_md.write_text(
        '## SLIDE 1 — template_index: 0 (Section Header)\n'
        '- placeholder: "Section Name Here" → "Intro"\n',
        encoding="utf-8",
    )
    plan = ContentPlanner.parse_slides_markdown(slides_md)
    assert plan.slides[0].diagram_path == ""
    assert plan.slides[0].intent == SlideIntent.BULLETS


# ---------------------------------------------------------------------------
# PPTX engine: diagram PNG embedding (image path, no drawio call)
# ---------------------------------------------------------------------------

def test_engine_embeds_png_diagram(tmp_path):
    """A slide with diagram_path pointing to a PNG gets an image shape added."""
    from pptx import Presentation
    from pptx.util import Inches
    from src.core.engines.pptx.engine import PPTXEngine
    from src.core.models import SlidePlan

    # Create a minimal 1-slide template
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    prs.slides.add_slide(blank)
    tpl = tmp_path / "tpl.pptx"
    prs.save(str(tpl))

    # Create a tiny real PNG (1x1 px)
    png = tmp_path / "diagram.png"
    # Minimal valid 1x1 white PNG bytes
    import struct, zlib
    def _png_1x1():
        sig = b'\x89PNG\r\n\x1a\n'
        def chunk(name, data):
            c = struct.pack('>I', len(data)) + name + data
            return c + struct.pack('>I', zlib.crc32(name + data) & 0xffffffff)
        ihdr = chunk(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0))
        raw = b'\x00\xff\xff\xff'
        idat = chunk(b'IDAT', zlib.compress(raw))
        iend = chunk(b'IEND', b'')
        return sig + ihdr + idat + iend
    png.write_bytes(_png_1x1())

    sdef = SlideDefinition(
        slide_number=1,
        template_index=0,
        slide_type="Diagram",
        intent=SlideIntent.DIAGRAM,
        diagram_path=str(png),
    )
    plan = SlidePlan(title="test", slides=[sdef])

    out = tmp_path / "out.pptx"
    engine = PPTXEngine()
    engine.render(plan, tpl, out)

    result = Presentation(str(out))
    # The rendered slide (index 0, after template slides removed) should have a picture shape
    shapes_with_pic = [s for s in result.slides[0].shapes if s.shape_type == 13]  # 13 = PICTURE
    assert len(shapes_with_pic) >= 1
