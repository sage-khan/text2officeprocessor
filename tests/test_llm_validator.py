"""
Tests for the LLMValidator semantic coherence check.
"""
import sys
from pathlib import Path
from unittest.mock import MagicMock

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from src.core.validation.validator import (
    LLMValidator,
    _extract_xlsx_text,
    _parse_llm_response,
)


# ---------------------------------------------------------------------------
# _parse_llm_response
# ---------------------------------------------------------------------------

def test_parse_clean_json_array():
    response = '[{"severity": "warning", "location": "Slide 1", "issue_type": "PLACEHOLDER_LEAK", "message": "Still contains template text"}]'
    issues = _parse_llm_response(response)
    assert len(issues) == 1
    assert issues[0]["issue_type"] == "PLACEHOLDER_LEAK"


def test_parse_empty_array():
    assert _parse_llm_response("[]") == []


def test_parse_markdown_fenced_json():
    response = "```json\n[{\"severity\": \"error\", \"location\": \"Slide 2\", \"issue_type\": \"TRUNCATED\", \"message\": \"Cut off mid\"}]\n```"
    issues = _parse_llm_response(response)
    assert len(issues) == 1
    assert issues[0]["severity"] == "error"


def test_parse_malformed_returns_empty():
    assert _parse_llm_response("not json at all") == []
    assert _parse_llm_response("{}") == []


def test_parse_embedded_array_in_prose():
    response = 'Here are the issues I found:\n[{"severity": "warning", "location": "Section 3", "issue_type": "EMPTY_SECTION", "message": "No body text"}]\nEnd.'
    issues = _parse_llm_response(response)
    assert len(issues) == 1
    assert issues[0]["issue_type"] == "EMPTY_SECTION"


# ---------------------------------------------------------------------------
# LLMValidator.validate — provider mocked
# ---------------------------------------------------------------------------

def _make_provider(response: str):
    p = MagicMock()
    p.generate.return_value = response
    return p


def test_validate_xlsx_no_issues(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Name", "Score"])
    ws.append(["Alice", 95])
    out = tmp_path / "report.xlsx"
    wb.save(str(out))

    provider = _make_provider("[]")
    validator = LLMValidator(provider=provider)
    result = validator.validate(out)

    assert result.passed
    assert result.issues == []
    provider.generate.assert_called_once()


def test_validate_xlsx_with_issues(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Section Name Here"])  # unreplaced placeholder
    out = tmp_path / "report.xlsx"
    wb.save(str(out))

    provider = _make_provider(
        '[{"severity": "warning", "location": "Sheet: Data", "issue_type": "PLACEHOLDER_LEAK", "message": "Template text found"}]'
    )
    validator = LLMValidator(provider=provider)
    result = validator.validate(out)

    assert len(result.issues) == 1
    assert "PLACEHOLDER_LEAK" in result.issues[0].message


def test_validate_gracefully_handles_llm_failure(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Some content"])
    out = tmp_path / "report.xlsx"
    wb.save(str(out))

    provider = MagicMock()
    provider.generate.side_effect = ConnectionError("LLM unreachable")

    validator = LLMValidator(provider=provider)
    result = validator.validate(out)

    # Should return empty result — never raises
    assert result.passed
    assert result.issues == []


def test_validate_unsupported_format_returns_empty(tmp_path):
    f = tmp_path / "file.csv"
    f.write_text("a,b,c\n1,2,3\n", encoding="utf-8")

    provider = _make_provider("[]")
    validator = LLMValidator(provider=provider)
    result = validator.validate(f)

    assert result.passed
    provider.generate.assert_not_called()


def test_validate_missing_file_returns_empty(tmp_path):
    provider = _make_provider("[]")
    validator = LLMValidator(provider=provider)
    result = validator.validate(tmp_path / "nonexistent.xlsx")

    assert result.passed
    provider.generate.assert_not_called()


# ---------------------------------------------------------------------------
# CLI --llm-validate flag wired correctly (no LLM, warns and skips)
# ---------------------------------------------------------------------------

def test_cli_llm_validate_without_llm_warns(tmp_path):
    from typer.testing import CliRunner
    from src.cli.main import app

    md = tmp_path / "test.md"
    md.write_text("# Title\n\n## Section\n\nContent here.\n", encoding="utf-8")
    out = tmp_path / "report.xlsx"

    runner = CliRunner()
    result = runner.invoke(app, [
        "convert",
        "--input", str(md),
        "--output", str(out),
        "--type", "xlsx",
        "--llm-validate",
        "--no-validate",
    ])

    assert result.exit_code == 0, result.output
    assert "WARN" in result.output or "Skipping LLM" in result.output
