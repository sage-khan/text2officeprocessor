"""
Tests for the PPTX Engine.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from src.core.engines.pptx.engine import PPTXEngine, sanitize_presentation, set_bullets
from src.core.exceptions import TemplateNotFoundError, RenderError
from src.core.planner.content_planner import ContentPlanner

TEMPLATE_PATH = Path(__file__).parent / "templates" / "sample-sections.pptx"
SLIDES_MD_PATH = Path(__file__).parent / "data" / "sample-slides.md"
MIN_EXPECTED_SIZE = 500_000  # 500KB minimum for a real background-preserved output


@pytest.fixture
def engine():
    return PPTXEngine()


@pytest.fixture
def slide_plan():
    if not SLIDES_MD_PATH.exists():
        pytest.skip("Test data not available")
    return ContentPlanner.parse_slides_markdown(SLIDES_MD_PATH)


def test_missing_template_raises_error(engine, slide_plan, tmp_path):
    with pytest.raises(TemplateNotFoundError):
        engine.render(slide_plan, tmp_path / "ghost.pptx", tmp_path / "out.pptx")


def test_empty_plan_raises_error(engine, tmp_path):
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    from src.core.models import SlidePlan
    empty_plan = SlidePlan(title="empty")
    with pytest.raises(RenderError, match="no slides"):
        engine.render(empty_plan, TEMPLATE_PATH, tmp_path / "out.pptx")


def test_render_produces_output_file(engine, slide_plan, tmp_path):
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    out = tmp_path / "output.pptx"
    result = engine.render(slide_plan, TEMPLATE_PATH, out)
    assert result.exists()


def test_output_is_valid_pptx(engine, slide_plan, tmp_path):
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    out = tmp_path / "output.pptx"
    engine.render(slide_plan, TEMPLATE_PATH, out)
    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == len(slide_plan.slides)


def test_output_preserves_backgrounds(engine, slide_plan, tmp_path):
    """File size heuristic: full backgrounds preserved means >500KB for 22 slides."""
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    out = tmp_path / "output.pptx"
    engine.render(slide_plan, TEMPLATE_PATH, out)
    assert out.stat().st_size >= MIN_EXPECTED_SIZE, (
        f"Output file {out.stat().st_size} bytes is suspiciously small — "
        "backgrounds may not be preserved."
    )


def test_no_markdown_artifacts_in_output(engine, slide_plan, tmp_path):
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    out = tmp_path / "output.pptx"
    engine.render(slide_plan, TEMPLATE_PATH, out)
    from pptx import Presentation
    prs = Presentation(str(out))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        for artifact in ["***", "**", "__"]:
                            assert artifact not in run.text, (
                                f"Markdown artifact '{artifact}' found in slide text: '{run.text[:60]}'"
                            )


def test_section_header_title_replaced(engine, slide_plan, tmp_path):
    """First slide should have the section title, not the template placeholder."""
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    out = tmp_path / "output.pptx"
    engine.render(slide_plan, TEMPLATE_PATH, out)
    from pptx import Presentation
    prs = Presentation(str(out))
    first_slide = prs.slides[0]
    all_text = " ".join(
        run.text
        for shape in first_slide.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
        if run.text.strip()
    )
    # The template placeholder must no longer appear — replaced with actual content
    assert "Section Name Here" not in all_text


def test_validator_passes_on_good_output(engine, slide_plan, tmp_path):
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    out = tmp_path / "output.pptx"
    engine.render(slide_plan, TEMPLATE_PATH, out)
    from src.core.validation.validator import ProgrammaticValidator
    result = ProgrammaticValidator().validate_pptx(out)
    assert result.passed
    errors = [i for i in result.issues if i.severity == "error"]
    assert len(errors) == 0


def test_fit_text_to_shape_never_shrinks_more_than_cap_below_original(tmp_path):
    """Regression test: overflow auto-shrink must never take a run more than
    MAX_SHRINK_FROM_ORIGINAL_PT below whatever size the template originally
    used for it, even if the absolute MIN_FONT_SIZE_PT floor would otherwise
    allow further reduction. A template's sizes are a deliberate design
    choice — past a point, cutting content is the right call, not shrinking
    a 28pt heading down toward 8pt to force an over-long paragraph to fit.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from src.core.engines.pptx.engine import (
        MAX_SHRINK_FROM_ORIGINAL_PT,
        MIN_FONT_SIZE_PT,
        fit_text_to_shape_single,
    )

    from pptx.enum.text import MSO_AUTO_SIZE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    tf = box.text_frame
    # A fresh textbox defaults to SHAPE_TO_FIT_TEXT, which `fit_text_to_shape_single`
    # deliberately skips (the template already handles that shape natively) --
    # disable it so this test exercises the manual shrink-loop path, the same
    # as a real fixed-size template placeholder would.
    tf.auto_size = MSO_AUTO_SIZE.NONE
    run = tf.paragraphs[0].add_run()
    original_size = 20
    run.font.size = Pt(original_size)
    # Deliberately huge overflowing text so the loop would run to completion
    # (hit `max_iterations`) rather than converge on its own — this is the
    # scenario where an uncapped loop would walk the size all the way down to
    # MIN_FONT_SIZE_PT.
    run.text = "Overflowing filler text. " * 200

    fit_text_to_shape_single(box)

    final_size = run.font.size.pt
    expected_floor = max(MIN_FONT_SIZE_PT, original_size - MAX_SHRINK_FROM_ORIGINAL_PT)
    assert final_size >= expected_floor, (
        f"run shrank to {final_size}pt, more than {MAX_SHRINK_FROM_ORIGINAL_PT}pt "
        f"below its original {original_size}pt (floor should be {expected_floor}pt)"
    )
    # And it actually WAS engaged (proves the test scenario really overflowed
    # and the cap is what stopped it, not a no-op).
    assert final_size < original_size


def test_duplicate_slide_preserves_relationship_ids():
    """Regression test for a real, confirmed defect: cloning a slide whose XML
    has more than one relationship (e.g. a picture placeholder's image rel
    plus the slide's slideLayout rel) previously used `get_or_add()` to copy
    relationships onto the new slide part, which mints a FRESH rId for each
    one rather than preserving the original. The deep-copied slide XML still
    has its original hardcoded `r:embed="rIdN"` reference, so if the new
    part's relationships get renumbered, that reference now points at the
    WRONG relationship (e.g. the slideLayout rel instead of the image rel) --
    the picture then fails to resolve and silently falls back to its
    solid-fill placeholder colour, rendering as a plain grey panel instead of
    the real image. `duplicate_slide()` must preserve each relationship's
    exact original rId.
    """
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    from pptx import Presentation
    from src.core.engines.pptx.engine import duplicate_slide

    prs = Presentation(str(TEMPLATE_PATH))

    # Find a bank slide with 2+ relationships (a picture placeholder plus its
    # slideLayout rel is the real-world case that exposed this bug).
    source_index = None
    for i, slide in enumerate(prs.slides):
        if len(slide.part.rels) >= 2:
            source_index = i
            break
    if source_index is None:
        pytest.skip("no multi-relationship bank slide in the shared test template")

    original_rel_map = {
        rid: (rel.reltype, str(rel.target_partname) if not rel.is_external else rel.target_ref)
        for rid, rel in prs.slides[source_index].part.rels.items()
    }

    new_slide = duplicate_slide(prs, source_index)
    new_rel_map = {
        rid: (rel.reltype, str(rel.target_partname) if not rel.is_external else rel.target_ref)
        for rid, rel in new_slide.part.rels.items()
    }

    assert new_rel_map == original_rel_map, (
        "duplicate_slide() must preserve every relationship's exact original "
        "rId -> (reltype, target); a renumbered rId breaks any hardcoded "
        "r:embed reference still present in the cloned slide XML.\n"
        f"original: {original_rel_map}\ngot:      {new_rel_map}"
    )


def test_set_bullets_removes_unused_paragraphs():
    """Regression test for the "phantom bullet" bug: a multi-paragraph bullet
    placeholder (the common "Rectangle 4"-style Multi Point template pattern)
    has more paragraph slots than the caller supplies bullets for. Unused
    paragraphs must be REMOVED from the XML entirely, not just have their run
    text blanked — an empty <a:p> can still carry an explicit <a:buChar>
    bullet-glyph definition in its <a:pPr> (inherited from the template's list
    style), which some renderers draw even with no text, producing a bullet
    marker with nothing next to it. This was a real, reported defect.
    """
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    from pptx import Presentation
    from src.core.engines.pptx.engine import duplicate_slide, set_bullets

    prs = Presentation(str(TEMPLATE_PATH))
    # Slide index 3 in the shared test template is the Multi Point bank slide:
    # a single "Rectangle 4" text frame with 12 mostly-whitespace paragraphs,
    # several of which carry an explicit bullet-glyph pPr (buChar/buNone).
    slide = duplicate_slide(prs, 3)
    ok = set_bullets(slide, ["Only two bullets here", "Second bullet, rest removed"])
    assert ok

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        paras = shape.text_frame.paragraphs
        if len(paras) < 2:
            continue
        texts = [p.text for p in paras]
        if texts[:2] != ["Only two bullets here", "Second bullet, rest removed"]:
            continue
        # This is the bullet box `set_bullets` filled — it must have EXACTLY
        # two paragraphs left, not two filled ones followed by leftover blanks.
        assert len(paras) == 2, (
            f"expected the unused paragraph slots to be removed, found "
            f"{len(paras)} paragraphs total: {texts!r}"
        )
        return
    pytest.fail("could not locate the bullet box that set_bullets filled")


def test_set_bullets_normalizes_indent_across_paragraph_slots():
    """Regression test for a real, reported defect: a bullet-box's paragraph SLOTS
    are not uniformly styled in the template XML — some slots have no marL/indent at
    all while others have a proper hanging indent for their bullet glyph. Filling
    bullets by position without normalizing indent meant whichever bullet landed in an
    un-indented slot wrapped flush-left while its siblings hang-indented, producing
    visibly inconsistent ("scattered") wrapping between bullets that are otherwise
    identically styled. set_bullets() must force every filled paragraph's marL/indent
    (and bullet glyph) to match paragraph 0's, regardless of what slot it landed in.
    """
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    from pptx import Presentation
    from src.core.engines.pptx.engine import duplicate_slide, set_bullets

    prs = Presentation(str(TEMPLATE_PATH))
    # Slide index 3 is the Multi Point bank slide (12 paragraph slots); fill all 12
    # so every slot — including any with no inherited marL/indent — gets used.
    slide = duplicate_slide(prs, 3)
    texts = [f"Bullet number {i}" for i in range(1, 13)]
    ok = set_bullets(slide, texts)
    assert ok

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        paras = shape.text_frame.paragraphs
        if [p.text for p in paras] != texts:
            continue
        ref_pPr = paras[0]._p.find(f"{A}pPr")
        ref_marL = ref_pPr.get("marL") if ref_pPr is not None else None
        ref_indent = ref_pPr.get("indent") if ref_pPr is not None else None
        for i, p in enumerate(paras):
            pPr = p._p.find(f"{A}pPr")
            marL = pPr.get("marL") if pPr is not None else None
            indent = pPr.get("indent") if pPr is not None else None
            assert marL == ref_marL, f"paragraph {i}: marL {marL!r} != paragraph 0's {ref_marL!r}"
            assert indent == ref_indent, f"paragraph {i}: indent {indent!r} != paragraph 0's {ref_indent!r}"
        return
    pytest.fail("could not locate the bullet box that set_bullets filled")


def test_set_big_statement_centers_text_and_removes_bullet_glyph():
    """set_big_statement() must turn a small one-bullet placeholder into a single,
    bold, centred, bullet-free paragraph with the other paragraph slots removed
    entirely (same phantom-glyph reasoning as set_bullets()) — not just place the
    text in the first paragraph and leave the box's small default size/position.
    """
    if not TEMPLATE_PATH.exists():
        pytest.skip("Template not available")
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR
    from src.core.engines.pptx.engine import duplicate_slide, set_big_statement

    prs = Presentation(str(TEMPLATE_PATH))
    # Slide index 2 in the shared test template is the Single Point bank slide:
    # a "Rectangle 4" text frame with 4 mostly-whitespace paragraphs.
    slide = duplicate_slide(prs, 2)
    statement = "A single statement that should dominate the slide, not hide in a corner"
    ok = set_big_statement(slide, statement)
    assert ok

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        paras = shape.text_frame.paragraphs
        if len(paras) != 1 or paras[0].text != statement:
            continue
        run = paras[0].runs[0]
        assert run.font.bold is True
        assert shape.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        return
    pytest.fail("could not locate the placeholder that set_big_statement filled")
