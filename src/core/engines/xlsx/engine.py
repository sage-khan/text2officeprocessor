"""
XLSX Engine — structured mapping via openpyxl.

Converts a SpreadsheetPlan into a formatted .xlsx workbook.
Each SheetDefinition becomes a worksheet with:
- Auto-detected column widths
- Bold header row
- Grid borders
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from src.core.exceptions import RenderError
from src.core.models import SpreadsheetPlan

logger = logging.getLogger(__name__)


class XLSXEngine:
    """
    Renders a SpreadsheetPlan into a .xlsx file using openpyxl.

    Usage:
        engine = XLSXEngine()
        engine.render(plan, output_path)
    """

    _INVALID_SHEET_CHARS = re.compile(r"[\[\]\*:/\\?]")

    @classmethod
    def _sanitize_sheet_name(cls, raw_name: str, used_names: set[str]) -> str:
        """Return an Excel-safe, unique worksheet title (max 31 chars)."""
        cleaned = cls._INVALID_SHEET_CHARS.sub("_", (raw_name or "").strip())
        cleaned = cleaned.strip("'")
        base = (cleaned[:31] or "Sheet").strip() or "Sheet"

        candidate = base
        i = 1
        while candidate in used_names:
            suffix = f"_{i}"
            candidate = f"{base[:31 - len(suffix)]}{suffix}"
            i += 1
        used_names.add(candidate)
        return candidate

    def render(self, plan: SpreadsheetPlan, output_path: Path) -> Path:
        """
        Generate a .xlsx workbook from a SpreadsheetPlan.

        Args:
            plan: The SpreadsheetPlan with all sheet definitions.
            output_path: Where to write the output .xlsx.

        Returns:
            The output_path on success.

        Raises:
            RenderError: If rendering fails.
        """
        output_path = Path(output_path)

        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RenderError("openpyxl is not installed. Run: pip install openpyxl") from exc

        if not plan.sheets:
            raise RenderError("SpreadsheetPlan contains no sheets to render.")

        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)  # Remove the default empty sheet

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        thin_border_side = Side(style="thin")
        cell_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side,
        )

        used_sheet_names: set[str] = set()
        for sheet_def in plan.sheets:
            sheet_name = self._sanitize_sheet_name(sheet_def.name, used_sheet_names)
            ws = workbook.create_sheet(title=sheet_name)

            # Header row
            if sheet_def.columns:
                for col_idx, header in enumerate(sheet_def.columns, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=str(header))
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = cell_border

            # Data rows
            row_start = 2 if sheet_def.columns else 1
            for row_idx, data_row in enumerate(sheet_def.rows, start=row_start):
                for col_idx, value in enumerate(data_row, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else "")
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.border = cell_border

            # Auto-fit column widths (estimate based on content length)
            for col_idx in range(1, max(len(sheet_def.columns), 1) + 1):
                max_length = 0
                col_letter = get_column_letter(col_idx)
                for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
                    for cell in row:
                        try:
                            cell_length = len(str(cell.value or ""))
                            if cell_length > max_length:
                                max_length = cell_length
                        except Exception:
                            pass
                adjusted_width = min(max_length + 4, 60)
                ws.column_dimensions[col_letter].width = adjusted_width

            # Freeze header row
            if sheet_def.columns:
                ws.freeze_panes = "A2"

            logger.debug(
                "Sheet '%s': %d column(s), %d row(s)",
                sheet_name,
                len(sheet_def.columns),
                len(sheet_def.rows),
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)

        logger.info("Saving XLSX output: %s", output_path.name)
        try:
            workbook.save(str(output_path))
        except Exception as exc:
            raise RenderError(f"Failed to save XLSX: {exc}") from exc

        return output_path
