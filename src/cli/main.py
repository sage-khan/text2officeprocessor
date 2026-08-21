"""
Text2OfficeProcessor CLI — entry point for the text2officeprocessor command-line tool.

Usage:
    text2officeprocessor convert --input input.md --template template.pptx --output output.pptx --type pptx
    text2officeprocessor convert --input input.md --template template.docx --output output.docx --type docx
    text2officeprocessor convert --input input.md --output output.xlsx --type xlsx
    text2officeprocessor convert --slides-md slides.md --template template.pptx --output out.pptx --type pptx
    text2officeprocessor convert --slides-md slides.md --template template.pptx --output out.pptx --config my-rules.yaml
    text2officeprocessor convert --config job.yaml   # input/template/output/type/llm settings all from one file
    text2officeprocessor analyze template.pptx
    text2officeprocessor batch --input-dir ./content/ --output-dir ./outputs/ --type pptx --template template.pptx
    text2officeprocessor batch --input-dir ./content/ --output-dir ./outputs/ --type xlsx
    text2officeprocessor drawio-export diagram.drawio --output diagram.png
    text2officeprocessor serve [--host 0.0.0.0] [--port 8000]
"""

from __future__ import annotations

import logging
import sys
from pathlib import Path
from typing import Any, Optional

import typer

from src.core.engines.docx.engine import DOCXEngine
from src.core.engines.pptx.engine import PPTXEngine
from src.core.engines.xlsx.engine import XLSXEngine
from src.core.exceptions import Text2OfficeProcessorError
from src.core.job_config import load_job_config, pick
from src.core.llm.runtime_config import resolve_provider_selection
from src.core.llm.providers import build_provider
from src.core.models import OutputFormat
from src.core.parser.preprocessor import InputPreprocessor
from src.core.planner.content_planner import ContentPlanner
from src.core.validation.validator import ProgrammaticValidator

app = typer.Typer(
    name="text2officeprocessor",
    help="Convert markdown / text / HTML to PPTX, DOCX, or XLSX using template-driven rendering.",
    add_completion=False,
)

# ---------------------------------------------------------------------------
# Bundled template resolution
# Supports both editable installs (src/data/templates/) and installed packages
# (via importlib.resources).
# ---------------------------------------------------------------------------

def _bundled_templates_dir() -> Path:
    """Return the directory containing bundled templates, regardless of install method."""
    try:
        from importlib.resources import files
        return Path(str(files("src.data").joinpath("templates")))
    except Exception:
        return Path(__file__).parent.parent / "data" / "templates"


def _resolve_template(template: Optional[Path], output_type: "OutputFormat") -> Optional[Path]:
    """Return the template path, falling back to the bundled generic template."""
    if template:
        return template
    tdir = _bundled_templates_dir()
    pptx = tdir / "generic-slides.pptx"
    docx = tdir / "generic-document.docx"
    if output_type.value == "pptx" and pptx.exists():
        return pptx
    if output_type.value == "docx" and docx.exists():
        return docx
    return None

LOG_LEVELS = {"debug": logging.DEBUG, "info": logging.INFO, "warning": logging.WARNING}


def _setup_logging(level: str) -> None:
    logging.basicConfig(
        format="%(asctime)s  %(levelname)-8s  %(name)s — %(message)s",
        level=LOG_LEVELS.get(level.lower(), logging.INFO),
    )


@app.command("convert")
def convert(
    input_file: Optional[Path] = typer.Option(
        None, "--input", "-i", help="Input .md / .txt / .html file."
    ),
    slides_md: Optional[Path] = typer.Option(
        None, "--slides-md", "-s",
        help="Pre-authored slides markdown file (## SLIDE N format). Bypasses LLM normalization.",
    ),
    template: Optional[Path] = typer.Option(
        None, "--template", "-t",
        help="Template .pptx or .docx file. Omit to use the bundled generic template.",
    ),
    output: Optional[Path] = typer.Option(
        None, "--output", "-o",
        help="Output file path. Required unless supplied via --config's job.output.",
    ),
    output_type: Optional[OutputFormat] = typer.Option(
        None, "--type", help="Output format: pptx | docx | xlsx. [default: pptx]"
    ),
    llm_provider: Optional[str] = typer.Option(
        None, "--llm", help="LLM provider: ollama | vllm | openai | claude | openrouter | groq | none."
    ),
    llm_model: Optional[str] = typer.Option(
        None, "--llm-model", help="Model name for the LLM provider."
    ),
    validate: Optional[bool] = typer.Option(
        None, "--validate/--no-validate", help="Run validation after rendering. [default: validate]"
    ),
    llm_validate: Optional[bool] = typer.Option(
        None, "--llm-validate/--no-llm-validate",
        help="Run an LLM semantic coherence check after rendering (requires --llm).",
    ),
    config: Optional[Path] = typer.Option(
        None, "--config", "-c",
        help=(
            "Path to a unified job config YAML. Can supply input/slides-md/template/output/type, "
            "LLM provider settings (including inline api_key or a local endpoint's base_url), and "
            "validate/overflow-strategy flags under `job:`/`llm:` keys — replacing most other flags. "
            "May also carry the rules sections (validation/sanitization/placeholder_map/...) normally "
            "read from config/default_rules.yaml. Explicit CLI flags always override this file."
        ),
    ),
    overflow_strategy: Optional[str] = typer.Option(
        None, "--overflow-strategy",
        help="PPTX overflow handling: summarize (LLM condense) | split (multi-slide) | shrink (font reduce) | auto (LLM decides).",
    ),
    log_level: Optional[str] = typer.Option(
        None, "--log-level", help="Logging level: debug | info | warning. [default: info]"
    ),
) -> None:
    """Convert an input document to PPTX, DOCX, or XLSX."""
    job_cfg = load_job_config(config)

    input_file = pick(input_file, job_cfg.input)
    slides_md = pick(slides_md, job_cfg.slides_md)
    template = pick(template, job_cfg.template)
    output = pick(output, job_cfg.output)
    llm_provider = pick(llm_provider, job_cfg.llm_provider)
    llm_model = pick(llm_model, job_cfg.llm_model)
    validate = pick(validate, job_cfg.validate)
    if validate is None:
        validate = True
    llm_validate = pick(llm_validate, job_cfg.llm_validate)
    if llm_validate is None:
        llm_validate = False
    overflow_strategy = pick(overflow_strategy, job_cfg.overflow_strategy)
    log_level = pick(log_level, job_cfg.log_level) or "info"
    output_type = OutputFormat(pick(output_type.value if output_type else None, job_cfg.output_type) or "pptx")

    if output is None:
        typer.echo("[ERROR] --output is required (either as a flag or via --config's job.output).", err=True)
        raise typer.Exit(code=1)

    _setup_logging(log_level)
    logger = logging.getLogger("text2officeprocessor.cli")

    # -- Resolve LLM provider (local-first; configurable)
    provider = None
    resolved_llm_name, llm_config = resolve_provider_selection(
        llm_provider, llm_model, job_cfg.llm_settings or None
    )
    if resolved_llm_name:
        try:
            provider = build_provider(resolved_llm_name, llm_config)
            logger.info("Using LLM provider: %s", resolved_llm_name)
        except Exception as exc:
            logger.warning(
                "Could not initialize LLM provider '%s': %s — proceeding without LLM.",
                resolved_llm_name,
                exc,
            )

    resolved_template = _resolve_template(template, output_type)
    if resolved_template and resolved_template != template:
        typer.echo(f"  Using bundled template: {resolved_template.name}")

    llm_validator = None
    if llm_validate and provider:
        from src.core.validation.validator import LLMValidator
        llm_validator = LLMValidator(provider=provider)
    elif llm_validate and not provider:
        typer.echo("  [WARN] --llm-validate requires --llm to be set. Skipping LLM validation.")

    try:
        if output_type == OutputFormat.PPTX:
            _run_pptx(input_file, slides_md, resolved_template, output, provider, validate, config, logger, llm_validator, overflow_strategy)
        elif output_type == OutputFormat.DOCX:
            _run_docx(input_file, resolved_template, output, provider, validate, config, logger, llm_validator)
        elif output_type == OutputFormat.XLSX:
            _run_xlsx(input_file, output, provider, validate, logger, llm_validator)
    except Text2OfficeProcessorError as exc:
        typer.echo(f"\n[ERROR] {exc}", err=True)
        raise typer.Exit(code=1)
    except Exception as exc:
        logger.exception("Unexpected error during conversion")
        typer.echo(f"\n[ERROR] Unexpected error: {exc}", err=True)
        raise typer.Exit(code=1)


def _run_pptx(
    input_file: Optional[Path],
    slides_md: Optional[Path],
    template: Optional[Path],
    output: Path,
    provider,
    validate: bool,
    config: Optional[Path],
    logger: logging.Logger,
    llm_validator=None,
    overflow_strategy: Optional[str] = None,
) -> None:
    from src.core.llm.normalizer import LLMNormalizer

    if not template:
        typer.echo("[ERROR] --template is required for PPTX output.", err=True)
        raise typer.Exit(code=1)

    if slides_md:
        # Direct mode: parse pre-authored slides markdown
        logger.info("Using pre-authored slides markdown: %s", slides_md)
        plan = ContentPlanner.parse_slides_markdown(slides_md)
    elif input_file:
        # Full pipeline mode
        logger.info("Running full pipeline on: %s", input_file)
        preprocessor = InputPreprocessor()
        parsed = preprocessor.parse(input_file)

        normalizer = LLMNormalizer(provider=provider)
        normalized = normalizer.normalize(parsed)

        planner = ContentPlanner(config_path=config)
        plan = planner.plan_slides(parsed, normalized)
    else:
        typer.echo("[ERROR] Either --input or --slides-md is required.", err=True)
        raise typer.Exit(code=1)

    typer.echo(f"  Slides to render: {len(plan.slides)}")

    # Content fitting for overflow handling
    engine = PPTXEngine(
        overflow_strategy=overflow_strategy,
        llm_provider=provider,
    )
    result_path = engine.render(plan, template, output)
    typer.echo(f"  Output written: {result_path}")

    if validate:
        validator = ProgrammaticValidator()
        validation = validator.validate_pptx(result_path)
        validator.print_report(validation)

    if llm_validator:
        from src.core.validation.validator import LLMValidator
        llm_result = llm_validator.validate(result_path)
        typer.echo("  LLM semantic validation:")
        ProgrammaticValidator().print_report(llm_result)


def _run_docx(
    input_file: Optional[Path],
    template: Optional[Path],
    output: Path,
    provider,
    validate: bool,
    config: Optional[Path],
    logger: logging.Logger,
    llm_validator=None,
) -> None:
    if not input_file:
        typer.echo("[ERROR] --input is required for DOCX output.", err=True)
        raise typer.Exit(code=1)
    if not template:
        typer.echo("[ERROR] --template is required for DOCX output.", err=True)
        raise typer.Exit(code=1)

    preprocessor = InputPreprocessor()
    parsed = preprocessor.parse(input_file)
    plan = ContentPlanner.plan_document(parsed)

    typer.echo(f"  Sections to render: {len(plan.sections)}")
    engine = DOCXEngine()
    result_path = engine.render(plan, template, output)
    typer.echo(f"  Output written: {result_path}")

    if validate:
        validator = ProgrammaticValidator()
        validation = validator.validate_docx(result_path)
        validator.print_report(validation)

    if llm_validator:
        llm_result = llm_validator.validate(result_path)
        typer.echo("  LLM semantic validation:")
        ProgrammaticValidator().print_report(llm_result)


def _run_xlsx(
    input_file: Optional[Path],
    output: Path,
    provider,
    validate: bool,
    logger: logging.Logger,
    llm_validator=None,
) -> None:
    if not input_file:
        typer.echo("[ERROR] --input is required for XLSX output.", err=True)
        raise typer.Exit(code=1)

    preprocessor = InputPreprocessor()
    parsed = preprocessor.parse(input_file)

    # Use LLM reorganizer if provider available for better sheet consolidation
    plan = ContentPlanner.plan_spreadsheet(parsed, llm_provider=provider)

    typer.echo(f"  Sheets to render: {len(plan.sheets)}")
    engine = XLSXEngine()
    result_path = engine.render(plan, output)
    typer.echo(f"  Output written: {result_path}")

    if validate:
        validator = ProgrammaticValidator()
        validation = validator.validate_xlsx(result_path)
        validator.print_report(validation)

    if llm_validator:
        llm_result = llm_validator.validate(result_path)
        typer.echo("  LLM semantic validation:")
        ProgrammaticValidator().print_report(llm_result)


def _classify_slide_structure(slide: Any) -> dict[str, Any]:
    """
    Inspect a template slide's shapes and produce structural signals an LLM
    (or human) can use to decide which markdown slide_type belongs here.

    Thin wrapper around `src.core.analysis.structural.classify_pptx_slide_structure`
    (the same structural pass the layout-manifest system uses) so the two
    never drift apart; renames its `guessed_content_affinity` key to this
    command's pre-existing `guessed_slide_type` for output compatibility.
    """
    from src.core.analysis.structural import classify_pptx_slide_structure

    signals = classify_pptx_slide_structure(slide)
    signals["guessed_slide_type"] = signals.pop("guessed_content_affinity")
    return signals


@app.command("analyze")
def analyze(
    template: Path = typer.Argument(..., help="Template .pptx file to analyze."),
    log_level: str = typer.Option("info", "--log-level", help="Logging level."),
) -> None:
    """Analyze a PPTX template: list every slide's full shape text plus a
    structural manifest so an LLM can decide which slide_type maps to which
    template_index for THIS template (the bundled DEFAULT_TEMPLATE_MAP only
    applies to the bundled generic template bank)."""
    _setup_logging(log_level)

    if not template.exists():
        typer.echo(f"[ERROR] Template not found: {template}", err=True)
        raise typer.Exit(code=1)

    try:
        from pptx import Presentation
    except ImportError:
        typer.echo("[ERROR] python-pptx is not installed.", err=True)
        raise typer.Exit(code=1)

    prs = Presentation(str(template))
    typer.echo(f"\nTemplate: {template.name}")
    typer.echo(f"Total slides: {len(prs.slides)}\n")

    manifest: list[dict[str, Any]] = []

    for i, slide in enumerate(prs.slides):
        layout_name = slide.slide_layout.name
        typer.echo(f"=== SLIDE {i} (layout: {layout_name}) ===")
        for j, shape in enumerate(slide.shapes):
            if not shape.has_text_frame:
                continue
            full_text = shape.text_frame.text
            if not full_text.strip():
                continue
            run_count = sum(len(p.runs) for p in shape.text_frame.paragraphs)
            multi_run = " [spans multiple runs/paragraphs]" if run_count > 1 else ""
            typer.echo(f'  Shape {j} "{shape.name}": {repr(full_text)}{multi_run}')
        signals = _classify_slide_structure(slide)
        typer.echo(
            f"  → structure: {signals['text_shape_count']} text shape(s), "
            f"{signals['picture_count']} picture(s), "
            f"bullet-group={signals['bullet_group_shapes']}, "
            f"pattern-A bullets={signals['pattern_a_bullet_paragraphs']}, "
            f"short text shapes={signals['short_text_shapes']}"
        )
        typer.echo(f"  → guessed slide_type: {signals['guessed_slide_type']} (advisory — verify against shape text above)")
        typer.echo("")
        manifest.append({
            "template_index": i,
            "layout": layout_name,
            **signals,
        })

    typer.echo("=== SLIDE-TYPE MANIFEST (for LLM slide-placement decisions) ===")
    typer.echo(
        "Use this table to pick a template_index for each markdown slide based on "
        "its content type. The 'guess' column is heuristic — cross-check it against "
        "the full shape text printed above before relying on it.\n"
    )
    typer.echo(f"{'idx':<5}{'layout':<16}{'guess':<32}{'bullets':<10}{'cards':<8}{'pics':<6}")
    for entry in manifest:
        typer.echo(
            f"{entry['template_index']:<5}"
            f"{entry['layout'][:14]:<16}"
            f"{entry['guessed_slide_type'][:30]:<32}"
            f"{max(entry['bullet_group_shapes'], entry['pattern_a_bullet_paragraphs']):<10}"
            f"{entry['short_text_shapes']:<8}"
            f"{entry['picture_count']:<6}"
        )
    typer.echo("")


@app.command("batch")
def batch(
    input_dir: Optional[Path] = typer.Option(
        None, "--input-dir", "-i",
        help="Directory containing input files (.md / .txt / .html). Required unless set via --config's job.input_dir.",
    ),
    output_dir: Optional[Path] = typer.Option(
        None, "--output-dir", "-o",
        help="Directory where output files will be written (created if absent). "
             "Required unless set via --config's job.output_dir.",
    ),
    output_type: Optional[OutputFormat] = typer.Option(
        None, "--type", help="Output format: pptx | docx | xlsx. [default: pptx]"
    ),
    template: Optional[Path] = typer.Option(
        None, "--template", "-t",
        help="Template .pptx or .docx file. Omit to use the bundled generic template.",
    ),
    glob_pattern: Optional[str] = typer.Option(
        None, "--pattern", "-p",
        help="Glob pattern to filter input files, e.g. '*.md' or 'section-*.html'. [default: *]",
    ),
    llm_provider: Optional[str] = typer.Option(
        None, "--llm", help="LLM provider: ollama | vllm | openai | claude | openrouter | groq | none."
    ),
    llm_model: Optional[str] = typer.Option(
        None, "--llm-model", help="Model name for the LLM provider."
    ),
    validate: Optional[bool] = typer.Option(
        None, "--validate/--no-validate", help="Run validation after each render. [default: validate]"
    ),
    llm_validate: Optional[bool] = typer.Option(
        None, "--llm-validate/--no-llm-validate",
        help="Run an LLM semantic coherence check after each render (requires --llm).",
    ),
    config: Optional[Path] = typer.Option(
        None, "--config", "-c",
        help=(
            "Path to a unified job config YAML — input_dir/output_dir/template/type/pattern, LLM "
            "provider settings, and validate flags under `job:`/`llm:` keys, plus optional rules "
            "sections. Explicit CLI flags always override this file."
        ),
    ),
    fail_fast: Optional[bool] = typer.Option(
        None, "--fail-fast/--no-fail-fast",
        help="Stop immediately on first error instead of continuing with remaining files.",
    ),
    log_level: Optional[str] = typer.Option(
        None, "--log-level", help="Logging level: debug | info | warning. [default: info]"
    ),
) -> None:
    """Convert every input file in a directory to the chosen output format."""
    job_cfg = load_job_config(config)

    input_dir = pick(input_dir, job_cfg.input_dir)
    output_dir = pick(output_dir, job_cfg.output_dir)
    template = pick(template, job_cfg.template)
    glob_pattern = pick(glob_pattern, job_cfg.pattern) or "*"
    llm_provider = pick(llm_provider, job_cfg.llm_provider)
    llm_model = pick(llm_model, job_cfg.llm_model)
    validate = pick(validate, job_cfg.validate)
    if validate is None:
        validate = True
    llm_validate = pick(llm_validate, job_cfg.llm_validate)
    if llm_validate is None:
        llm_validate = False
    fail_fast = pick(fail_fast, job_cfg.fail_fast)
    if fail_fast is None:
        fail_fast = False
    log_level = pick(log_level, job_cfg.log_level) or "info"
    output_type = OutputFormat(pick(output_type.value if output_type else None, job_cfg.output_type) or "pptx")

    _setup_logging(log_level)
    logger = logging.getLogger("text2officeprocessor.batch")

    if input_dir is None or output_dir is None:
        typer.echo(
            "[ERROR] --input-dir and --output-dir are required "
            "(either as flags or via --config's job.input_dir/job.output_dir).",
            err=True,
        )
        raise typer.Exit(code=1)

    if not input_dir.is_dir():
        typer.echo(f"[ERROR] Input directory not found: {input_dir}", err=True)
        raise typer.Exit(code=1)

    output_dir.mkdir(parents=True, exist_ok=True)

    # Collect candidate files matching the pattern and supported formats
    supported_suffixes = {".md", ".txt", ".html", ".htm"}
    candidates = sorted(
        f for f in input_dir.glob(glob_pattern)
        if f.is_file() and f.suffix.lower() in supported_suffixes
    )

    if not candidates:
        typer.echo(
            f"[WARN] No supported files found in '{input_dir}' matching '{glob_pattern}'. "
            "Supported extensions: .md .txt .html .htm"
        )
        raise typer.Exit(code=0)

    # Resolve LLM provider once for the whole batch (local-first; configurable)
    provider = None
    resolved_llm_name, llm_config = resolve_provider_selection(
        llm_provider, llm_model, job_cfg.llm_settings or None
    )
    if resolved_llm_name:
        try:
            provider = build_provider(resolved_llm_name, llm_config)
            logger.info("Using LLM provider: %s", resolved_llm_name)
        except Exception as exc:
            logger.warning(
                "Could not initialize LLM provider '%s': %s — proceeding without LLM.",
                resolved_llm_name,
                exc,
            )

    resolved_template = _resolve_template(template, output_type)
    if resolved_template and resolved_template != template:
        typer.echo(f"  Using bundled template: {resolved_template.name}")

    llm_validator = None
    if llm_validate and provider:
        from src.core.validation.validator import LLMValidator
        llm_validator = LLMValidator(provider=provider)
    elif llm_validate and not provider:
        typer.echo("  [WARN] --llm-validate requires --llm to be set. Skipping LLM validation.")

    ext_map = {OutputFormat.PPTX: ".pptx", OutputFormat.DOCX: ".docx", OutputFormat.XLSX: ".xlsx"}
    out_ext = ext_map[output_type]

    total = len(candidates)
    succeeded: list[Path] = []
    failed: list[tuple[Path, str]] = []

    typer.echo(f"\nBatch: {total} file(s) → {output_type.value.upper()} in '{output_dir}'\n")

    for idx, input_file in enumerate(candidates, start=1):
        output_file = output_dir / (input_file.stem + out_ext)
        typer.echo(f"  [{idx}/{total}] {input_file.name} → {output_file.name}")
        try:
            if output_type == OutputFormat.PPTX:
                _run_pptx(input_file, None, resolved_template, output_file, provider, validate, config, logger, llm_validator)
            elif output_type == OutputFormat.DOCX:
                _run_docx(input_file, resolved_template, output_file, provider, validate, config, logger, llm_validator)
            elif output_type == OutputFormat.XLSX:
                _run_xlsx(input_file, output_file, provider, validate, logger, llm_validator)
            succeeded.append(output_file)
        except Exception as exc:
            msg = str(exc)
            failed.append((input_file, msg))
            typer.echo(f"    [FAILED] {msg}", err=True)
            if fail_fast:
                typer.echo("\n[ABORTED] --fail-fast is set. Stopping batch.", err=True)
                raise typer.Exit(code=1)

    # Summary
    typer.echo(f"\n{'='*50}")
    typer.echo(f"Batch complete: {len(succeeded)}/{total} succeeded, {len(failed)} failed.")
    if failed:
        typer.echo("\nFailed files:")
        for path, reason in failed:
            typer.echo(f"  {path.name}: {reason}")
        raise typer.Exit(code=1)


@app.command("watch")
def watch(
    input_file: Optional[Path] = typer.Option(
        None, "--input", "-i",
        help="Input .md / .txt / .html file to watch. Required unless set via --config's job.input.",
    ),
    output: Optional[Path] = typer.Option(
        None, "--output", "-o", help="Output file path. Required unless set via --config's job.output.",
    ),
    output_type: Optional[OutputFormat] = typer.Option(
        None, "--type", help="Output format: pptx | docx | xlsx. [default: pptx]"
    ),
    template: Optional[Path] = typer.Option(
        None, "--template", "-t",
        help="Template .pptx or .docx file. Omit to use the bundled generic template.",
    ),
    slides_md: Optional[Path] = typer.Option(
        None, "--slides-md", "-s",
        help="Pre-authored slides markdown. Watched alongside --input when provided.",
    ),
    llm_provider: Optional[str] = typer.Option(
        None, "--llm", help="LLM provider for normalization."
    ),
    llm_model: Optional[str] = typer.Option(
        None, "--llm-model", help="Model name for the LLM provider."
    ),
    validate: Optional[bool] = typer.Option(
        None, "--validate/--no-validate", help="Run validation after each regeneration. [default: validate]"
    ),
    config: Optional[Path] = typer.Option(
        None, "--config", "-c",
        help="Path to a unified job config YAML (input/template/output/type/llm settings/validate), "
             "plus optional rules sections. Explicit CLI flags always override this file.",
    ),
    debounce: Optional[float] = typer.Option(
        None, "--debounce", help="Seconds to wait after a change before regenerating. [default: 1.0]"
    ),
    log_level: Optional[str] = typer.Option(
        None, "--log-level", help="Logging level: debug | info | warning. [default: info]"
    ),
) -> None:
    """Watch an input file and auto-regenerate the output on every save."""
    job_cfg = load_job_config(config)

    input_file = pick(input_file, job_cfg.input)
    output = pick(output, job_cfg.output)
    template = pick(template, job_cfg.template)
    slides_md = pick(slides_md, job_cfg.slides_md)
    llm_provider = pick(llm_provider, job_cfg.llm_provider)
    llm_model = pick(llm_model, job_cfg.llm_model)
    validate = pick(validate, job_cfg.validate)
    if validate is None:
        validate = True
    debounce = pick(debounce, job_cfg.debounce)
    if debounce is None:
        debounce = 1.0
    log_level = pick(log_level, job_cfg.log_level) or "info"
    output_type = OutputFormat(pick(output_type.value if output_type else None, job_cfg.output_type) or "pptx")

    _setup_logging(log_level)
    logger = logging.getLogger("text2officeprocessor.watch")

    if input_file is None or output is None:
        typer.echo(
            "[ERROR] --input and --output are required (either as flags or via --config's job.input/job.output).",
            err=True,
        )
        raise typer.Exit(code=1)

    if not input_file.exists():
        typer.echo(f"[ERROR] Input file not found: {input_file}", err=True)
        raise typer.Exit(code=1)

    try:
        from watchdog.events import FileSystemEventHandler
        from watchdog.observers import Observer
    except ImportError:
        typer.echo("[ERROR] watchdog is not installed. Run: pip install watchdog", err=True)
        raise typer.Exit(code=1)

    provider = None
    resolved_llm_name, llm_cfg = resolve_provider_selection(llm_provider, llm_model, job_cfg.llm_settings or None)
    if resolved_llm_name:
        try:
            provider = build_provider(resolved_llm_name, llm_cfg)
            logger.info("Using LLM provider: %s", resolved_llm_name)
        except Exception as exc:
            logger.warning("Could not initialize LLM provider '%s': %s", resolved_llm_name, exc)

    resolved_template = _resolve_template(template, output_type)
    if resolved_template and resolved_template != template:
        typer.echo(f"  Using bundled template: {resolved_template.name}")

    # Files to monitor (deduped)
    watched_files = {input_file.resolve()}
    if slides_md:
        watched_files.add(slides_md.resolve())

    import threading
    import time

    _timer: list = [None]
    _lock = threading.Lock()

    def _regenerate() -> None:
        typer.echo(f"\n[{__import__('datetime').datetime.now().strftime('%H:%M:%S')}] Change detected — regenerating...")
        try:
            if output_type == OutputFormat.PPTX:
                _run_pptx(input_file, slides_md, resolved_template, output, provider, validate, config, logger)
            elif output_type == OutputFormat.DOCX:
                _run_docx(input_file, resolved_template, output, provider, validate, config, logger)
            elif output_type == OutputFormat.XLSX:
                _run_xlsx(input_file, output, provider, validate, logger)
            typer.echo("  Done.")
        except Exception as exc:
            typer.echo(f"  [ERROR] {exc}", err=True)

    def _schedule_regenerate() -> None:
        with _lock:
            if _timer[0] is not None:
                _timer[0].cancel()
            _timer[0] = threading.Timer(debounce, _regenerate)
            _timer[0].start()

    class _ChangeHandler(FileSystemEventHandler):
        def on_modified(self, event):
            if Path(event.src_path).resolve() in watched_files:
                _schedule_regenerate()

        def on_created(self, event):
            if Path(event.src_path).resolve() in watched_files:
                _schedule_regenerate()

    # Run once immediately on start
    _regenerate()

    watch_dirs = {p.parent for p in watched_files}
    observer = Observer()
    handler = _ChangeHandler()
    for d in watch_dirs:
        observer.schedule(handler, str(d), recursive=False)

    observer.start()
    file_list = ", ".join(p.name for p in watched_files)
    typer.echo(f"\nWatching: {file_list}  (Ctrl+C to stop)\n")

    try:
        while True:
            time.sleep(0.5)
    except KeyboardInterrupt:
        observer.stop()
        typer.echo("\nWatch mode stopped.")
    finally:
        observer.join()


@app.command("drawio-export")
def drawio_export(
    input_file: Path = typer.Argument(..., help="Source .drawio file to export."),
    output: Optional[Path] = typer.Option(
        None, "--output", "-o",
        help="Output PNG path. Defaults to <input>.png alongside the source file.",
    ),
    scale: str = typer.Option("2", "--scale", "-s", help="Export scale factor (default: 2 for 2x resolution)."),
    border: str = typer.Option("10", "--border", "-b", help="Border width in pixels around the diagram."),
    transparent: bool = typer.Option(False, "--transparent/--no-transparent", help="Transparent PNG background."),
    page: Optional[int] = typer.Option(None, "--page", "-p", help="1-based page index (default: first page)."),
    all_pages: bool = typer.Option(False, "--all-pages", help="Export every page as separate PNGs."),
    log_level: str = typer.Option("info", "--log-level", help="Logging level: debug | info | warning."),
) -> None:
    """Export a .drawio diagram to PNG using the drawio CLI."""
    _setup_logging(log_level)

    from src.core.drawio.converter import (
        DrawioExportError,
        export_all_pages,
        export_drawio_to_png,
    )

    if not input_file.exists():
        typer.echo(f"[ERROR] File not found: {input_file}", err=True)
        raise typer.Exit(code=1)

    try:
        if all_pages:
            out_dir = output if output and output.is_dir() else input_file.parent
            paths = export_all_pages(input_file, output_dir=out_dir, scale=scale, border=border)
            for p in paths:
                typer.echo(f"  Exported: {p}")
        else:
            path = export_drawio_to_png(
                input_file,
                output_path=output,
                page_index=page,
                scale=scale,
                border=border,
                transparent=transparent,
            )
            typer.echo(f"  Exported: {path}")
    except DrawioExportError as exc:
        typer.echo(f"\n[ERROR] {exc}", err=True)
        raise typer.Exit(code=1)
    except Exception as exc:
        typer.echo(f"\n[ERROR] Unexpected error: {exc}", err=True)
        raise typer.Exit(code=1)


@app.command("templates")
def list_templates() -> None:
    """List the bundled generic templates included with text2officeprocessor."""
    typer.echo("\nBundled templates (use with --template or omit for auto-selection):\n")

    found_any = False
    for path in sorted(_bundled_templates_dir().glob("*")):
        if path.suffix not in (".pptx", ".docx"):
            continue
        found_any = True
        size_kb = path.stat().st_size // 1024
        if path.suffix == ".pptx":
            try:
                from pptx import Presentation
                prs = Presentation(str(path))
                detail = f"{len(prs.slides)} slides"
            except Exception:
                detail = "PPTX"
        else:
            detail = "DOCX"
        typer.echo(f"  {path.name:<35} {detail:<15} {size_kb} KB")
        typer.echo(f"  Path: {path}\n")

    if not found_any:
        typer.echo("  No bundled templates found. Run: python scripts/create_bundled_templates.py")


@app.command("extract-styles")
def extract_styles_cmd(
    template: Path = typer.Argument(..., exists=True, readable=True, help="Path to .docx or .pptx template."),
    output: Path = typer.Option(..., "--output", "-o", help="Output JSON file path."),
) -> None:
    """Extract the complete visual identity from a template into a JSON style sheet."""
    from src.core.extraction.style_extractor import extract_styles
    try:
        stylesheet = extract_styles(template)
        stylesheet.save(output)
        style_count = len(stylesheet.styles) if stylesheet.styles else len(stylesheet.slide_layouts)
        typer.echo(f"  Extracted {style_count} styles/layouts → {output}")
    except Exception as exc:
        typer.echo(f"[ERROR] {exc}", err=True)
        raise typer.Exit(code=1)


@app.command("extract")
def extract_cmd(
    source: Path = typer.Argument(..., exists=True, readable=True, help="Path to .docx, .pptx, or .xlsx file."),
    output_dir: Path = typer.Option(..., "--output", "-o", help="Output directory for extracted files."),
    no_styles: bool = typer.Option(False, "--no-styles", help="Skip style extraction."),
    no_media: bool = typer.Option(False, "--no-media", help="Skip media extraction."),
) -> None:
    """Extract content, styles, and media from an Office file into Markdown."""
    from src.core.extraction.content_extractor import extract_content
    try:
        content_path = extract_content(
            source, output_dir,
            extract_styles=not no_styles,
            extract_media=not no_media,
        )
        typer.echo(f"  Content → {content_path}")
        if not no_styles:
            suffix = source.suffix.lower()
            styles_name = "pptx_styles.json" if suffix == ".pptx" else "styles.json"
            sp = output_dir / styles_name
            if sp.exists():
                typer.echo(f"  Styles  → {sp}")
        media_dir = output_dir / "media"
        if media_dir.exists() and any(media_dir.iterdir()):
            count = sum(1 for _ in media_dir.iterdir())
            typer.echo(f"  Media   → {media_dir}/ ({count} files)")
    except Exception as exc:
        typer.echo(f"[ERROR] {exc}", err=True)
        raise typer.Exit(code=1)


@app.command("diff-styles")
def diff_styles_cmd(
    before: Path = typer.Argument(..., exists=True, readable=True, help="First styles.json file."),
    after: Path = typer.Argument(..., exists=True, readable=True, help="Second styles.json file."),
) -> None:
    """Compare two extracted style sheets and show differences."""
    import json as _json
    try:
        a = _json.loads(before.read_text(encoding="utf-8"))
        b = _json.loads(after.read_text(encoding="utf-8"))

        a_styles = {s["style_id"]: s for s in a.get("styles", [])}
        b_styles = {s["style_id"]: s for s in b.get("styles", [])}

        added = set(b_styles) - set(a_styles)
        removed = set(a_styles) - set(b_styles)
        common = set(a_styles) & set(b_styles)

        changes_found = False
        for sid in sorted(added):
            typer.echo(f"ADDED    {sid} ({b_styles[sid].get('name', '')})")
            changes_found = True
        for sid in sorted(removed):
            typer.echo(f"REMOVED  {sid} ({a_styles[sid].get('name', '')})")
            changes_found = True
        for sid in sorted(common):
            if a_styles[sid] != b_styles[sid]:
                typer.echo(f"CHANGED  {sid}")
                _diff_dict(a_styles[sid], b_styles[sid], prefix="  ")
                changes_found = True

        if not changes_found:
            typer.echo("No style differences found.")
    except Exception as exc:
        typer.echo(f"[ERROR] {exc}", err=True)
        raise typer.Exit(code=1)


def _diff_dict(a: dict, b: dict, prefix: str = "") -> None:
    """Print key-level differences between two dicts."""
    all_keys = sorted(set(list(a.keys()) + list(b.keys())))
    for key in all_keys:
        va, vb = a.get(key), b.get(key)
        if va != vb:
            typer.echo(f"{prefix}{key}: {va!r} → {vb!r}")


def _generate_manifest_cmd(
    template: Path, manifest_out: Path, llm_provider: Optional[str], llm_model: Optional[str],
) -> None:
    """Generate a layout manifest for `template` (structural pass, plus an
    LLM classification pass if a provider resolves and is reachable) and
    write it to exactly the path the caller asked for — deliberately calls
    `generate_pptx_manifest`/`generate_docx_manifest` directly rather than
    `get_or_generate_manifest()`, which also writes a second copy to the
    template's own directory (`<stem>.layout-manifest.json`, the caching
    convention `ContentPlanner` consults); a caller who names an explicit
    `--manifest` path shouldn't get a surprise second file next to their
    template."""
    from src.core.analysis.manifest import generate_docx_manifest, generate_pptx_manifest

    provider = None
    resolved_llm_name, llm_config = resolve_provider_selection(llm_provider, llm_model)
    if resolved_llm_name:
        try:
            provider = build_provider(resolved_llm_name, llm_config)
            typer.echo(f"  Using LLM provider for classification pass: {resolved_llm_name}")
        except Exception as exc:
            typer.echo(
                f"  [WARN] Could not initialize LLM provider '{resolved_llm_name}': {exc} "
                "— generating structural-only manifest.",
            )

    suffix = template.suffix.lower()
    try:
        if suffix == ".pptx":
            manifest = generate_pptx_manifest(template, llm_provider=provider)
        elif suffix == ".docx":
            manifest = generate_docx_manifest(template, llm_provider=provider)
        else:
            typer.echo(
                f"[ERROR] Layout manifests are not supported for '{suffix}' templates "
                "(PPTX and DOCX only — see docs/feature.md).",
                err=True,
            )
            raise typer.Exit(code=1)
    except typer.Exit:
        raise
    except Exception as exc:
        typer.echo(f"[ERROR] Manifest generation failed: {exc}", err=True)
        raise typer.Exit(code=1)

    manifest_out.parent.mkdir(parents=True, exist_ok=True)
    manifest_out.write_text(manifest.to_json(), encoding="utf-8")

    typer.echo(f"\n  Layout manifest written: {manifest_out}")
    typer.echo(f"  Format: {manifest.format}  classified_by_llm: {manifest.classified_by_llm}")
    if manifest.slides:
        typer.echo(f"  Slides analyzed: {len(manifest.slides)}")
        for s in manifest.slides:
            affinity = ", ".join(s.content_affinity) or "unknown"
            typer.echo(f"    [{s.index}] {s.layout_name or '(unnamed layout)'} — {affinity} — {len(s.slots)} slot(s)")
    if manifest.sections:
        typer.echo(f"  Sections analyzed: {len(manifest.sections)}")
        for sec in manifest.sections:
            typer.echo(f"    {sec.role} — table={sec.supports_table} image={sec.has_image_anchor}")
    typer.echo("")


@app.command("analyze-template")
def analyze_template_cmd(
    template: Path = typer.Argument(..., exists=True, readable=True, help="Path to .docx or .pptx template."),
    manifest_out: Optional[Path] = typer.Option(
        None, "--manifest",
        help="Generate a layout manifest (structural pass, PPTX/DOCX only) and write it as JSON here.",
    ),
    llm_provider: Optional[str] = typer.Option(
        None, "--llm", help="LLM provider for the manifest's classification pass: ollama | vllm | openai | claude | openrouter | groq | none."
    ),
    llm_model: Optional[str] = typer.Option(
        None, "--llm-model", help="Model name for the LLM provider."
    ),
) -> None:
    """Print a summary of a template's styles, layouts, and fonts. With
    --manifest, also generate a layout manifest (see docs/feature.md)."""
    if manifest_out is not None:
        _generate_manifest_cmd(template, manifest_out, llm_provider, llm_model)
        return

    from src.core.extraction.style_extractor import extract_styles
    try:
        ss = extract_styles(template)
        typer.echo(f"\n  Template: {ss.source_file}")
        typer.echo(f"  Format:   {ss.format.upper()}")
        typer.echo(f"  Theme:    major={ss.theme.major_font or 'n/a'}  minor={ss.theme.minor_font or 'n/a'}")
        typer.echo(f"  Colors:   {len(ss.theme.colors)} theme colors")

        if ss.format == "docx":
            typer.echo(f"  Styles:   {len(ss.styles)}")
            typer.echo(f"  Numbering: {len(ss.numbering_defs)} definitions")
            if ss.section_properties:
                sp = ss.section_properties
                typer.echo(f"  Page:     {sp.page_width_pt:.0f}x{sp.page_height_pt:.0f}pt ({sp.orientation})")
                typer.echo(f"  Margins:  L={sp.margin_left_pt:.0f} R={sp.margin_right_pt:.0f} T={sp.margin_top_pt:.0f} B={sp.margin_bottom_pt:.0f}")
            # Show key paragraph styles
            para_styles = [s for s in ss.styles if s.style_type == "paragraph"]
            typer.echo(f"\n  Paragraph styles ({len(para_styles)}):")
            for s in para_styles[:20]:
                font_info = ""
                if s.font:
                    parts = []
                    if s.font.name:
                        parts.append(s.font.name)
                    if s.font.size_pt:
                        parts.append(f"{s.font.size_pt}pt")
                    if s.font.bold:
                        parts.append("bold")
                    if s.font.color:
                        parts.append(s.font.color)
                    font_info = " — " + ", ".join(parts) if parts else ""
                typer.echo(f"    {s.style_id}: {s.name}{font_info}")

        elif ss.format == "pptx":
            typer.echo(f"  Slides:   {len(ss.slide_layouts)}")
            w_in = ss.slide_width_emu / 914400
            h_in = ss.slide_height_emu / 914400
            typer.echo(f"  Size:     {w_in:.1f}\" x {h_in:.1f}\"")
            typer.echo(f"\n  Slide layouts:")
            for layout in ss.slide_layouts:
                typer.echo(f"    [{layout.index}] {layout.name} — {len(layout.placeholders)} shapes")
                for ph in layout.placeholders[:5]:
                    typer.echo(f"        {ph.name}: {ph.type} ({ph.width_pt:.0f}x{ph.height_pt:.0f}pt)")

        typer.echo("")
    except Exception as exc:
        typer.echo(f"[ERROR] {exc}", err=True)
        raise typer.Exit(code=1)


@app.command("serve")
def serve(
    host: str = typer.Option("127.0.0.1", "--host", help="Bind host (default: 127.0.0.1)."),
    port: int = typer.Option(8000, "--port", "-p", help="Bind port (default: 8000)."),
    reload: bool = typer.Option(False, "--reload", help="Enable auto-reload (development only)."),
    log_level: str = typer.Option("info", "--log-level", help="Uvicorn log level."),
) -> None:
    """Start the Text2OfficeProcessor web UI server."""
    try:
        import uvicorn
    except ImportError:
        typer.echo(
            "[ERROR] Web UI requires extra dependencies. Install with:\n"
            "  pip install text2officeprocessor[web]\n"
            "or: pip install fastapi 'uvicorn[standard]' python-multipart",
            err=True,
        )
        raise typer.Exit(code=1)

    typer.echo(f"\nText2OfficeProcessor Web UI — http://{host}:{port}\nPress Ctrl+C to stop.\n")
    uvicorn.run(
        "src.web.app:create_app",
        factory=True,
        host=host,
        port=port,
        reload=reload,
        log_level=log_level.lower(),
    )


def main() -> None:
    app()


if __name__ == "__main__":
    main()
